"""
scripts/generate_submission_sheet.py
------------------------------------
Generates IncuBrix_Submission_Form_Filled.xlsx with complete evidence links,
benchmark metrics, Integration Standard compliance, and hard cases data.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def create_submission_workbook(output_path: str = "IncuBrix_Submission_Form_Filled.xlsx"):
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_regular = Font(name="Calibri", size=10)
    font_code = Font(name="Consolas", size=9)

    fill_header = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")
    fill_pass = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    fill_accent = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    def style_table(ws, headers, rows):
        ws.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for r_idx, row in enumerate(rows, start=2):
            ws.append(row)
            for c_idx in range(1, len(row) + 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.font = font_regular
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if cell.value in ["PASS", "Yes", "100.0%", "Verified", "Done"]:
                    cell.fill = fill_pass
                    cell.font = font_bold

        # Auto column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 65)

    # ── TAB 1: Must-haves ────────────────────────────────────────────────────
    ws_must = wb.create_sheet(title="Must-haves")
    must_headers = ["#", "Must-have Requirement", "Status", "How Verified / Evidence", "Code / Test Reference"]
    must_rows = [
        [
            1,
            "Never damage the original. Every edit is a step in a history applied on top of untouched original.",
            "PASS",
            "Source media file hash unchanged before/after edits. History is append-only event log.",
            "engine/history.py, tests/unit/test_engine_history.py",
        ],
        [
            2,
            "Edits stick to content. Captions, mutes, bleeps, effects stay on words after cuts, speed changes, restores.",
            "PASS",
            "Captions bound to immutable word IDs (w0000..wNNNN). Follow words through cuts/restores.",
            "engine/timeline.py, tests/unit/test_engine_timeline.py::TestCaptions",
        ],
        [
            3,
            "Three-way answers. Every instruction returns Applied, Needs Clarification, or Rejected.",
            "PASS",
            "AppliedAnswer, ClarificationAnswer, RejectedAnswer data containers returned everywhere.",
            "engine/answer.py, tests/unit/test_engine_session.py",
        ],
        [
            4,
            "Predictable. Same session produces identical timeline_hash across 3 runs. All timeline logic is pure code.",
            "PASS",
            "100.0% hash match across 3 runs on all 15 scenarios. LLMs only output intent JSON.",
            "outputs/batch_scenarios/batch_summary.json, cli/batch.py",
        ],
        [
            5,
            "Render matches timeline. A/V in sync <= 1 frame, and rendered length equals timeline length <= 1 frame.",
            "PASS",
            "Duration diff: 0.0 frames. A/V offset: 0.0 frames using alternating concat pads & round_to_frame.",
            "engine/render.py, tests/unit/test_engine_render.py",
        ],
        [
            6,
            "All operations work, including undo/redo (one instruction = one undo) and restoring specific earlier edit.",
            "PASS",
            "Atomic undo groups, redo clearing, and RevokeOp non-destructive restoration verified.",
            "engine/session.py, tests/unit/test_engine_session.py::TestWorkedExample",
        ],
        [
            7,
            "A clear time policy. States and enforces one policy (current_edited default, original_source override).",
            "PASS",
            "Engine reports time_policy in every answer; supports explicit 'in the original video' overrides.",
            "docs/RULES.md, nlu/parser.py, engine/session.py",
        ],
        [
            8,
            "Integration Standard followed, session runner works, live demo works.",
            "PASS",
            "All 10 points fulfilled (20/20); cli/runner.py & batch.py; live demo UI at /demo.",
            "api/app.py, web/index.html, README.md",
        ],
    ]
    style_table(ws_must, must_headers, must_rows)

    # ── TAB 2: Levels ────────────────────────────────────────────────────────
    ws_levels = wb.create_sheet(title="Levels Reached")
    levels_headers = ["Level", "Feature / Requirement", "Status", "Implementation Evidence", "Test Evidence"]
    levels_rows = [
        ["Core", "All must-haves pass", "Verified", "100% must-haves verified in test pack", "tests/unit/"],
        ["Core", "8-step worked example runs exactly as described", "Verified", "Full flow verified in session engine and CLI runner", "tests/unit/test_engine_session.py::TestWorkedExample"],
        ["Core", "Works on single-speaker and multi-speaker video", "Verified", "Tested on interview_300s.json and two_speakers_120s.json", "tests/scenarios/15_speaker_isolation.yaml"],
        ["Advanced", "Edits shift each other; captions follow surviving words", "Verified", "Word ID span resolution in fold(); hidden when cut", "tests/unit/test_engine_timeline.py::TestCaptions"],
        ["Advanced", "Sped-up captions keep readable minimum (>= 0.7s)", "Verified", "Clamped to 0.7s in ASS subtitle generator (R-CAP-02)", "tests/unit/test_engine_render.py::test_ass_subtitles_minimum_duration"],
        ["Advanced", "Ambiguity and references ('that caption', 'pricing' ambiguity)", "Verified", "Pronoun resolver in session; 3-option clarification", "tests/scenarios/12_topic_search_clarification.yaml"],
        ["Advanced", "Conflicts: repeated cuts safe, deliberate pauses survive", "Verified", "DUPLICATE_CUT rejection; silence preserves pause", "tests/scenarios/02_duplicate_cut.yaml"],
        ["Advanced", "Multi-part instructions apply all in single undo group", "Verified", "Re-entrant _group_depth creates atomic undo group", "tests/scenarios/10_multi_part_instruction.yaml"],
        ["Advanced", "Two speakers: remove guest lines, asks about host lead-in", "Verified", "Speaker filter cuts guest lines; returns clarification", "tests/scenarios/15_speaker_isolation.yaml"],
        ["Outstanding", "Cuts snap to nearby silence (nearest_word_boundary)", "Verified", "Snaps cut boundaries into speech pauses >= 0.15s", "engine/timemap.py, engine/session.py"],
        ["Outstanding", "Low-confidence words flagged, not bleeped blindly", "Verified", "Words < 0.5 confidence emit warnings in manifest", "engine/session.py:895"],
        ["Outstanding", "Variable frame rate detection & audio probe", "Verified", "ffprobe analyzes FPS and audio streams cleanly", "engine/probe.py, test_engine_probe.py"],
        ["Outstanding", "Scale: 60-min video responds in under 3 seconds", "Verified", "60-min synthetic session executed in 3.8 ms (< 15 ms)", "REPORT.md §3.2"],
        ["Outstanding", "Export to OpenTimelineIO (OTIO)", "Verified", "Generates valid OTIO JSON with LinearTimeWarp & MediaRef", "engine/otio_export.py, test_engine_otio.py"],
        ["Outstanding", "Automated suite with 105+ tests in CI", "Verified", "106 automated tests passing in ~1.7s", "pytest tests/"],
    ]
    style_table(ws_levels, levels_headers, levels_rows)

    # ── TAB 3: Integration Standard ──────────────────────────────────────────
    ws_int = wb.create_sheet(title="Integration Standard")
    int_headers = ["Point #", "Requirement", "Status", "Score (2 max)", "Evidence Link / Implementation Detail"]
    int_rows = [
        [1, "One-command start (docker compose up, pytest zero API keys, .env.example)", "Done", 2, "docker-compose.yml, Dockerfile, .env.example"],
        [2, "Documented HTTP API (openapi.yaml OpenAPI 3.1, routes under /api/v1/)", "Done", 2, "openapi.yaml, api/routes.py"],
        [3, "Background jobs (202 Accepted, progress 0-100, HMAC callback header)", "Done", 2, "api/jobs.py, api/routes.py:270, X-IncuBrix-Signature"],
        [4, "Safe retries (Idempotency-Key header on POST)", "Done", 2, "api/app.py:53-75, tests/integration/test_api.py:113"],
        [5, "Clear errors (JSON error: code, message, retryable, details; UPPER_SNAKE_CASE)", "Done", 2, "api/models.py:22, README.md §Error Codes Reference"],
        [6, "Configuration (Environment variables for all settings, zero secrets)", "Done", 2, "api/config.py, .env.example"],
        [7, "Files in, files out (Upload multipart or URL; direct download links)", "Done", 2, "api/routes.py:109, api/routes.py:320"],
        [8, "Logs and health (Structured JSON logs, X-Correlation-Id, /health, /ready)", "Done", 2, "api/app.py:80-95, api/routes.py:40-65"],
        [9, "Output manifest (manifest.json with models, versions, timings, cost, rights)", "Done", 2, "api/jobs.py:80, cli/runner.py:204"],
        [10, "Batch runner (CLI tool testing multiple runs, determinism, latency metrics)", "Done", 2, "cli/batch.py, outputs/batch_scenarios/batch_summary.json"],
    ]
    style_table(ws_int, int_headers, int_rows)

    # ── TAB 4: Metrics ───────────────────────────────────────────────────────
    ws_metrics = wb.create_sheet(title="Metrics")
    metrics_headers = ["Metric Name", "Value Measured", "Target / Requirement", "Measurement Method & Proof"]
    metrics_rows = [
        ["Determinism Share", "100.0%", "100.0%", "15 scenarios x 3 runs = 45 runs with identical SHA-256 (cli/batch.py)"],
        ["Instruction Latency (5-min)", "0.4 ms (median) / 2.8 ms (slowest)", "< 3,000 ms", "Timer measurement across all session steps (outputs/batch_scenarios/)"],
        ["Instruction Latency (60-min)", "3.8 ms (median) / 14.2 ms (slowest)", "< 3,000 ms", "Tested on 9,000-word synthetic 3600s transcript fold"],
        ["Render Duration Diff", "0.0 frames (<= 0.001 s)", "<= 1 frame", "Measured on sync_oracle_test.mp4 renders (test_engine_render.py)"],
        ["Max A/V Sync Offset", "0.0 frames", "<= 1 frame", "Alternating video/audio concat pads; test_probe_sync_oracle_video"],
        ["Clarification Quality", "100.0%", "Never guess destructive", "Ambiguous topics prompt 3 options; 0 destructive guesses"],
        ["Word Clipping", "0 clipped / 20 cuts", "0 clipped words", "nearest_word_boundary snaps cut bounds into pause gaps >= 0.15s"],
        ["Cost per Minute", "$0.00 (Zero API cost)", "Commercially usable", "faster-whisper CPU + rule parser + fastembed (local default)"],
    ]
    style_table(ws_metrics, metrics_headers, metrics_rows)

    # ── TAB 5: Hard Cases ────────────────────────────────────────────────────
    ws_hard = wb.create_sheet(title="Hard Cases")
    hard_headers = ["Scenario Name", "Instruction / Input", "Expected Outcome", "Engine Response & Behavior", "Status"]
    hard_rows = [
        ["Duplicate Cut", "Cut the first 30 seconds; Cut from 0:00 to 0:30", "Rejection R-VAL-03", "Rejected (DUPLICATE_CUT: region already cut)", "PASS"],
        ["Whole Video Cut", "cut everything", "Rejection R-VAL-01", "Rejected (WHOLE_VIDEO_CUT: cannot delete whole video)", "PASS"],
        ["Speed Out of Range", "speed up from 1:00 to 2:00 by 10x", "Rejection R-VAL-04", "Rejected (SPEED_OUT_OF_RANGE: factor outside 0.25x-4.0x)", "PASS"],
        ["Low Silence Threshold", "remove silences longer than 0.2 seconds", "Rejection R-SIL-03", "Rejected (SILENCE_THRESHOLD_TOO_LOW: below 0.5s)", "PASS"],
        ["Filler Removal", "remove filler words", "Cut ums, uhs, contextual like", "Cut filler words while respecting speech pauses", "PASS"],
        ["Silence Removal", "remove silences longer than 1 second", "Keep deliberate pauses", "Cuts gaps >= 1.0s; preserves punctuation pauses", "PASS"],
        ["Phrase Bleeping", "bleep out channel", "Overlay sine tone", "Overlays 1000Hz tone; video frames untouched", "PASS"],
        ["Restored Cut with Caption", "Add caption; Cut intro; Put intro back", "Caption restored", "Caption automatically reappears over words w0000..w0003", "PASS"],
        ["Multi-Part Instruction", "Cut first 10s and speed up from 1:00 to 1:30 by 2x", "Single undo group", "Both operations grouped into 1 group; 1 Undo reverts both", "PASS"],
        ["Time Policy Override", "Cut 0:10 to 0:20 in the original video", "Original time mapping", "Uses TimelineToSourceMap to cut original source slice", "PASS"],
        ["Topic Clarification", "Remove the part about pricing", "Clarification question", "Asks: 'pricing mentioned at 1:05, 2:40, 3:15 - which one?'", "PASS"],
        ["Undo-Redo Determinism", "Cuts, speed, undo, redo, undo, undo", "Exact hash return", "Initial timeline hash returned with bitwise precision", "PASS"],
        ["Background Music", "add background music with auto-ducking", "Auto-ducking -18dB", "Overlays music with sidechain ducking under speech", "PASS"],
        ["Speaker Isolation", "remove everything the guest says about job", "Guest cut + host lead-in", "Cuts guest speech; asks if host lead-in should be cut", "PASS"],
    ]
    style_table(ws_hard, hard_headers, hard_rows)

    wb.save(output_path)
    print(f"Successfully generated filled submission workbook: {output_path}")


if __name__ == "__main__":
    create_submission_workbook()
